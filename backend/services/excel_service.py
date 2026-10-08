"""
excel_service.py
Reads and writes Q&A records to an Excel workbook.
Each role gets its own worksheet. Uses openpyxl (sync, run in threadpool).

Column layout (v2 — Role column added at position 5):
  A  ID
  B  Session ID
  C  Timestamp
  D  Question
  E  Role            ← NEW
  F  Points (JSON)
  G  Illustrated Answer
  H  Source
"""

import json
from pathlib import Path
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill

DATA_DIR = Path(__file__).resolve().parents[2] / "data"
EXCEL_FILE = DATA_DIR / "interview_qa_database.xlsx"

ROLES = [
    "Test Manager",
    "Scrum Master",
    "Defect Manager",
    "SAP Tester",
    "SAP Test Manager",
    "Project Manager",
]

HEADER = ["ID", "Session ID", "Timestamp", "Question", "Role", "Points (JSON)", "Illustrated Answer", "Source"]
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")

#           A    B    C    D     E    F    G    H
COL_WIDTHS = [36, 36, 24, 60,   20,  80,  100, 10]


def _sheet_name(role: str) -> str:
    """Return a safe Excel sheet name (max 31 chars, no special chars)."""
    safe = "".join(c for c in role if c not in r'/\?*[]:"')
    return safe[:31]


def _get_or_create_workbook() -> openpyxl.Workbook:
    DATA_DIR.mkdir(parents=True, exist_ok=True)

    if EXCEL_FILE.exists():
        wb = openpyxl.load_workbook(EXCEL_FILE)
    else:
        wb = openpyxl.Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    for role in ROLES:
        name = _sheet_name(role)
        if name not in wb.sheetnames:
            ws = wb.create_sheet(title=name)
            ws.append(HEADER)
            for cell, width in zip(ws[1], COL_WIDTHS):
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                ws.column_dimensions[cell.column_letter].width = width

    return wb


def save_qa(record: dict) -> None:
    """Append one Q&A record to the correct role sheet."""
    wb = _get_or_create_workbook()
    name = _sheet_name(record["role"])
    ws = wb[name]

    ws.append([
        record.get("id", ""),
        record.get("session_id", ""),
        record.get("timestamp", datetime.utcnow().isoformat()),
        record.get("question", ""),
        record.get("role", ""),                        # col E — Role
        json.dumps(record.get("points", [])),           # col F — Points (JSON)
        record.get("illustrated", ""),                  # col G — Illustrated Answer
        record.get("source", "ai"),                     # col H — Source
    ])

    wb.save(EXCEL_FILE)


def get_history(role: str) -> list[dict]:
    """Return all saved Q&A rows for a given role."""
    if not EXCEL_FILE.exists():
        return []

    wb = _get_or_create_workbook()
    name = _sheet_name(role)
    if name not in wb.sheetnames:
        return []

    ws = wb[name]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[0]:          # skip empty rows
            continue
        try:
            points = json.loads(row[5]) if row[5] else []   # col F (index 5)
        except (json.JSONDecodeError, TypeError):
            points = []
        rows.append({
            "id":          row[0],
            "session_id":  row[1],
            "timestamp":   str(row[2]) if row[2] else "",
            "question":    row[3] or "",
            "role":        row[4] or role,               # col E (index 4)
            "points":      points,
            "illustrated": row[6] or "",                 # col G (index 6)
            "source":      row[7] or "ai",               # col H (index 7)
        })

    return rows
